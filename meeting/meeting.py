"""회의 녹음 -> 전사 -> 요약 -> 바탕화면 저장.

사용법: python meeting.py            (Enter 치면 녹음 종료 후 자동 처리)
        python meeting.py a.wav b.wav (이미 있는 녹음 파일로 회의록만 뽑기)

전사는 세그먼트마다 .txt.part에 적어두므로 중간에 끊겨도 다시 돌리면 이어서 한다.
"""

import datetime
import os
import re
import shutil
import sys
import tempfile
import threading
import time
import warnings
import wave
import winreg

import numpy as np
import soundcard as sc

warnings.filterwarnings("ignore", message="data discontinuity in recording")

SR = 48000
BLOCK = SR // 10
# ponytail: 정확도/속도 조절 손잡이. 속도 우선이면 base, 정확도 아쉬우면 small/medium, GPU면 large-v3.
#           바꾼 뒤 다시 뽑으려면 기존 .txt 전사본을 지워야 캐시를 안 쓴다.
#           base/small/medium은 로컬 캐시에 있다. 그 외 모델은 첫 실행에 인터넷이 필요하다.
WHISPER_MODEL = os.getenv("WHISPER_MODEL", "base")

# ponytail: 한 줄 유지. 여러 줄이면 CLI에 형식 지시가 안 붙는다.
PROMPT = (
    "입력된 업무 회의 녹취록을 한국어 회의록으로 정리하세요. "
    "반드시 이 5개 제목을 순서대로 그대로 사용하세요: "
    "'## 한 줄 요약', '## 핵심 논의', '## 결정사항', '## 액션 아이템', '## 후속 확인 필요'. "
    "액션 아이템은 '담당자 / 할 일 / 기한' 형태로 쓰고 언급이 없으면 '미지정'이라고 적으세요. "
    "녹취에 없는 내용은 지어내지 말고, 알아듣기 어려운 부분은 '(불명확)'으로 표시하세요. "
    "서두나 맺음말 없이 회의록 본문만 출력하세요."
)


def desktop():
    """OneDrive 리다이렉트를 반영한 실제 바탕화면 경로."""
    key = r"Software\Microsoft\Windows\CurrentVersion\Explorer\Shell Folders"
    with winreg.OpenKey(winreg.HKEY_CURRENT_USER, key) as k:
        return winreg.QueryValueEx(k, "Desktop")[0]


def mix(*tracks):
    """길이가 다른 모노 트랙들을 합쳐 클리핑 없이 정규화."""
    tracks = [t for t in tracks if len(t)]
    if not tracks:
        return np.zeros(0, "float32")
    n = max(len(t) for t in tracks)
    out = sum(np.pad(t, (0, n - len(t))) for t in tracks)
    peak = float(np.abs(out).max())
    return out / peak * 0.95 if peak > 1.0 else out


class Track(threading.Thread):
    """한 입력 장치를 Enter 칠 때까지 계속 읽어들이는 스레드."""

    def __init__(self, mic, label):
        super().__init__(daemon=True)
        self.mic, self.label = mic, label
        self.frames = []
        self.error = None
        self.stopping = threading.Event()

    def run(self):
        try:
            with self.mic.recorder(samplerate=SR, blocksize=BLOCK) as rec:
                while not self.stopping.is_set():
                    self.frames.append(rec.record(numframes=BLOCK))
        except Exception as e:
            self.error = e

    def mono(self):
        if not self.frames:
            return np.zeros(0, "float32")
        a = np.concatenate(self.frames)
        return (a.mean(axis=1) if a.ndim > 1 else a).astype("float32")


def open_tracks():
    """마이크(내 목소리) + 스피커 루프백(화상회의 상대방 목소리)."""
    tracks = []
    try:
        tracks.append(Track(sc.default_microphone(), "마이크"))
    except Exception as e:
        print(f"  마이크 없음: {e}")
    try:
        spk = sc.default_speaker().name
        tracks.append(Track(sc.get_microphone(spk, include_loopback=True), "시스템 소리"))
    except Exception as e:
        print(f"  시스템 소리 캡처 불가: {e}")
    return tracks


def write_wav(path, mono):
    with wave.open(path, "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(SR)
        w.writeframes((np.clip(mono, -1, 1) * 32767).astype("<i2").tobytes())


def resume_from(part_path):
    """중단된 전사의 (이미 받아쓴 줄들, 다시 시작할 초) 를 돌려준다."""
    lines, offset = [], 0.0
    if os.path.exists(part_path):
        with open(part_path, encoding="utf-8") as f:
            for line in f:
                end, _, text = line.rstrip("\n").partition("\t")
                if text:
                    lines.append(text)
                    offset = float(end)
    return lines, offset


def transcribe(wav_path):
    """전사. 완성본은 .txt로 캐시하고, 진행 중에는 .txt.part에 한 줄씩 적는다."""
    from faster_whisper import WhisperModel
    from faster_whisper.audio import decode_audio

    txt_path = os.path.splitext(wav_path)[0] + ".txt"
    if os.path.exists(txt_path):
        print("  기존 전사본 재사용")
        return open(txt_path, encoding="utf-8").read().strip()

    part_path = txt_path + ".part"
    lines, offset = resume_from(part_path)
    if offset:
        print(f"  이어서 전사: {offset / 60:.0f}분 지점부터")

    audio = decode_audio(wav_path)[int(offset * 16000):]   # decode_audio는 16kHz 모노
    model = WhisperModel(WHISPER_MODEL, device="cpu", compute_type="int8",
                         cpu_threads=os.cpu_count() or 4)
    # beam_size=1(그리디)이 빔서치보다 몇 배 빠르다. 정확도는 조금 손해.
    segments, _ = model.transcribe(audio, language="ko", vad_filter=True,
                                   beam_size=1, condition_on_previous_text=False)

    with open(part_path, "a", encoding="utf-8") as f:
        for s in segments:
            text = s.text.strip()
            if not text:
                continue
            lines.append(text)
            f.write(f"{offset + s.end:.2f}\t{text}\n")
            f.flush()
            print(f"\r  전사 {(offset + s.end) / 60:.0f}분째", end="", flush=True)
    print()

    out = "\n".join(lines)
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(out + "\n")
    os.remove(part_path)
    return out


def summarize(transcript):
    """Claude Code CLI로 요약. 기존 Pro 구독을 쓰므로 API 키가 필요 없다."""
    import subprocess

    exe = shutil.which("claude.cmd") or shutil.which("claude")
    if not exe:
        return "(claude CLI를 찾을 수 없어 요약을 건너뜁니다. 아래 전문을 참고하세요.)"
    r = subprocess.run(
        [exe, "-p", PROMPT, "--permission-mode", "bypassPermissions"], input=transcript, capture_output=True,
        text=True, encoding="utf-8", errors="replace", timeout=900,
    )
    if r.returncode != 0 or not r.stdout.strip():
        return f"(요약 실패: {r.stderr.strip()[:500] or '빈 응답'}. 아래 전문을 참고하세요.)"
    return r.stdout.strip()


def to_docx(md_path):
    """우리가 쓴 마크다운을 워드 문서로 변환. 반환값은 .docx 경로."""
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "맑은 고딕"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

    for raw in open(md_path, encoding="utf-8").read().splitlines():
        line = re.sub(r"\*\*(.+?)\*\*", r"\1", raw).rstrip()
        if not line or set(line) <= set("-|= "):   # 빈 줄, 구분선, 표 구분행
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], 2)
        elif line.startswith("## "):
            doc.add_heading(line[3:], 1)
        elif line.startswith("# "):
            doc.add_heading(line[2:], 0)
        elif line.startswith(("- ", "* ")):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.startswith("|"):
            doc.add_paragraph("   ".join(c.strip() for c in line.strip("|").split("|")))
        else:
            doc.add_paragraph(line)

    out = os.path.splitext(md_path)[0] + ".docx"
    doc.save(out)
    return out


def report(wav_path):
    """녹음 파일 하나 -> 전사 -> 요약 -> .md/.docx. 이미 있는 wav에도 그대로 쓴다."""
    wav_path = os.path.abspath(wav_path)
    base = os.path.splitext(os.path.basename(wav_path))[0]
    with wave.open(wav_path, "rb") as w:
        secs = w.getnframes() / w.getframerate()
    dur = f"{secs:.0f}초" if secs < 60 else f"{secs / 60:.0f}분"
    try:
        stamp = datetime.datetime.strptime(base, "회의록_%Y-%m-%d_%H%M")
    except ValueError:                  # 이름이 다르면 파일 수정시각을 회의 시각으로
        stamp = datetime.datetime.fromtimestamp(os.path.getmtime(wav_path))

    print(f"\n[{base}] {dur}. 전사 중… (회의 길이만큼 걸립니다)")
    transcript = transcribe(wav_path)
    if not transcript:
        sys.exit(f"전사 결과가 비었습니다. 녹음 파일: {wav_path}")

    print("요약 중…")
    summary = summarize(transcript)

    md_path = os.path.splitext(wav_path)[0] + ".md"
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(f"# 회의록 {stamp:%Y-%m-%d %H:%M} ({dur})\n\n")
        f.write(summary)
        f.write("\n\n---\n\n## 전체 녹취록\n\n" + transcript + "\n")

    try:
        out = to_docx(md_path)
    except Exception as e:              # 워드 변환이 실패해도 마크다운은 이미 남아 있다
        print(f"  워드 변환 실패: {e}")
        out = md_path
    print(f"✓ 문서 생성: {out}")
    return out


def main():
    stamp = datetime.datetime.now()
    tracks = open_tracks()
    for t in tracks:
        t.start()
    time.sleep(0.5)

    live = []
    for t in tracks:
        if t.error:
            print(f"  {t.label} 실패: {t.error}")
        else:
            live.append(t)
    if not live:
        sys.exit("녹음 장치를 하나도 못 열었습니다.")

    print(f"\n● 녹음 중 [{', '.join(t.label for t in live)}]. 회의 끝나면 Enter.")
    try:
        input()
    except (KeyboardInterrupt, EOFError):
        pass
    for t in live:
        t.stopping.set()
    for t in live:
        t.join(timeout=3)

    audio = mix(*(t.mono() for t in live))
    secs = len(audio) / SR
    if secs < 3:
        sys.exit("녹음된 소리가 없습니다.")
    dur = f"{secs:.0f}초" if secs < 60 else f"{secs / 60:.0f}분"

    # 바탕화면에는 결과 문서 하나만 남긴다. wav/txt/md 중간 파일은 임시 폴더에.
    work = os.path.join(tempfile.gettempdir(), "회의록")
    os.makedirs(work, exist_ok=True)
    wav_path = os.path.join(work, stamp.strftime("회의록_%Y-%m-%d_%H%M") + ".wav")
    write_wav(wav_path, audio)
    print(f"■ 녹음 종료 ({dur}).")

    out = report(wav_path)
    final = os.path.join(desktop(), os.path.basename(out))
    shutil.move(out, final)
    print(f"✓ 바탕화면: {os.path.basename(final)}   (원본 녹음: {work})")
    os.startfile(final)


def selftest():
    a = np.array([1.0, 1.0, 1.0], "float32")
    b = np.array([1.0], "float32")
    m = mix(a, b)
    assert len(m) == 3, m
    assert abs(m[0] - 0.95) < 1e-6, m       # 2.0 피크 -> 0.95로 정규화
    assert abs(m[1] - 0.475) < 1e-6, m      # 패딩된 뒷부분은 절반
    assert len(mix()) == 0
    q = np.array([0.5, -0.5], "float32")
    assert np.allclose(mix(q), q)           # 피크 <= 1 이면 그대로
    assert os.path.isdir(desktop()), desktop()

    part = os.path.join(os.environ["TEMP"], "_mt.txt.part")
    assert resume_from(part + ".nope") == ([], 0.0)          # 파일 없으면 처음부터
    with open(part, "w", encoding="utf-8") as f:
        f.write("12.50\t안녕하세요\n\n30.00\t네 그럼 시작\t하죠\n")
    assert resume_from(part) == (["안녕하세요", "네 그럼 시작\t하죠"], 30.0), resume_from(part)
    os.remove(part)

    md = os.path.join(os.environ["TEMP"], "_mt.md")
    with open(md, "w", encoding="utf-8") as f:
        f.write("# 제목\n\n## 절\n\n- 항목 **강조**\n\n---\n\n| a | b |\n|---|---|\n| 1 | 2 |\n\n본문\n")
    dx = to_docx(md)
    from docx import Document
    texts = [p.text for p in Document(dx).paragraphs]
    assert "제목" in texts and "절" in texts, texts
    assert "항목 강조" in texts, texts                    # ** 제거
    assert "a   b" in texts and "1   2" in texts, texts    # 표 행 보존
    assert "본문" in texts, texts
    assert not [t for t in texts if t and set(t) <= set("-| ")], texts  # 구분선 제거
    os.remove(md)
    os.remove(dx)

    tracks = open_tracks()
    assert tracks, "입력 장치를 하나도 못 찾음"
    for t in tracks:
        t.start()
    time.sleep(1.5)
    for t in tracks:
        t.stopping.set()
    for t in tracks:
        t.join(timeout=3)
        print(f"  {t.label}: {'실패 ' + str(t.error) if t.error else f'{len(t.mono())/SR:.2f}초'}")
    assert any(t.error is None and len(t.mono()) for t in tracks), "녹음된 샘플이 없음"
    print("ok")


if __name__ == "__main__":
    wavs = [a for a in sys.argv[1:] if not a.startswith("-")]
    if "--selftest" in sys.argv:
        selftest()
    elif wavs:
        for p in wavs:
            report(p)
    else:
        main()
